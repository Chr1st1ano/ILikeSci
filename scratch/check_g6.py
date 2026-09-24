import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('masterlists/GRADE-6-MASTERLIST.xlsx', data_only=True)
print("GRADE 6 Sheets:", wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    non_empty = [r for r in ws.iter_rows(values_only=True) if any(r)]
    print(f"Sheet '{s}': non-empty rows = {len(non_empty)}, first row = {non_empty[0] if non_empty else None}")
