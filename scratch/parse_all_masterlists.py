import openpyxl
import re
import sys
import unicodedata

sys.stdout.reconfigure(encoding='utf-8')

def clean_name(raw_name):
    if raw_name is None:
        return ""
    name = unicodedata.normalize('NFC', str(raw_name))
    name = re.sub(r'[\t\r\n]+', ' ', name)
    name = re.sub(r'\s+', ' ', name)
    name = name.strip()
    # Strip any trailing ", -"
    name = re.sub(r',\s*-\s*$', '', name).strip()
    return name

def parse_all():
    wb4 = openpyxl.load_workbook('masterlists/Grade-4-Masterlist-BCES-2026-2027.xlsx', data_only=True)
    g4_data = []
    
    for sname in wb4.sheetnames:
        ws = wb4[sname]
        # Clean section name: "G4 Maagap" -> "Maagap"
        section = sname.replace('G4', '').replace('g4', '').strip()
        current_gender = 'M'
        m_count = 0
        f_count = 0
        students = []
        for r in range(1, ws.max_row + 1):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            non_empty = [c for c in row if c is not None and str(c).strip() != '']
            if not non_empty:
                continue
            r_str = " ".join(str(c).strip() for c in non_empty).upper()
            if any(k in r_str for k in ['PREPARED BY', 'TEACHER', 'ADVISER', 'BCES', 'GRADE 4', 'G4']):
                continue
            # Check col 0 (number) and col 1 (name)
            val0 = str(row[0]).strip() if row[0] is not None else ""
            val1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            val2 = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            
            # Check gender flag in row or col
            if 'FEMALE' in r_str or 'GIRL' in r_str:
                current_gender = 'F'
                continue
            if 'MALE' in r_str or 'BOY' in r_str:
                current_gender = 'M'
                continue
                
            # If val0 is number '1' and we already had students, switch to female
            if val0.isdigit():
                num = int(val0)
                if num == 1 and len(students) > 0 and current_gender == 'M':
                    current_gender = 'F'
                
                # Check if val1 is a name
                if val1 and not any(k in val1.upper() for k in ['PREPARED', 'TEACHER', 'NAME', 'GRADE']):
                    gender = current_gender
                    if val2 in ['M', 'F']:
                        gender = val2
                    cname = clean_name(val1)
                    if cname:
                        students.append({
                            'grade': '4',
                            'section': section,
                            'sheet': sname,
                            'num': num,
                            'name': cname,
                            'gender': gender,
                            'row': r
                        })
                        if gender == 'M': m_count += 1
                        else: f_count += 1
            else:
                # Row didn't have number in col 0? Let's check if col 1 has number or if col 0 has name
                # Let's inspect anomalous rows
                if any(c.isdigit() for c in val0) or val1:
                    print(f"  [G4 Check Row {r} in {sname}]: {row}")

        g4_data.append((sname, section, m_count, f_count, students))

    print("=== GRADE 4 SUMMARY ===")
    total_g4 = 0
    for sname, sec, mc, fc, stus in g4_data:
        print(f"Section {sec} ({sname}): {mc} Males, {fc} Females -> Total: {len(stus)}")
        total_g4 += len(stus)
    print(f"TOTAL GRADE 4 STUDENTS: {total_g4}\n")

    # GRADE 6
    wb6 = openpyxl.load_workbook('masterlists/GRADE-6-MASTERLIST.xlsx', data_only=True)
    g6_data = []
    for sname in wb6.sheetnames:
        ws = wb6[sname]
        section = sname.replace('GRADE 6 -', '').replace('GRADE 6', '').strip()
        current_gender = 'M'
        m_count = 0
        f_count = 0
        students = []
        for r in range(1, ws.max_row + 1):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            non_empty = [c for c in row if c is not None and str(c).strip() != '']
            if not non_empty:
                continue
            r_str = " ".join(str(c).strip() for c in non_empty).upper()
            if any(k in r_str for k in ['PREPARED BY', 'TEACHER', 'ADVISER', 'BCES', 'GRADE 6', 'G6', 'SY 202', 'NAME OF LEARNER']):
                continue
            
            val0 = str(row[0]).strip() if row[0] is not None else ""
            val1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            
            if 'FEMALE' in r_str:
                current_gender = 'F'
                continue
            if 'MALE' in r_str:
                current_gender = 'M'
                continue
                
            if val0.isdigit():
                num = int(val0)
                if val1 and not any(k in val1.upper() for k in ['PREPARED', 'TEACHER', 'NAME', 'GRADE']):
                    cname = clean_name(val1)
                    if cname:
                        students.append({
                            'grade': '6',
                            'section': section,
                            'sheet': sname,
                            'num': num,
                            'name': cname,
                            'gender': current_gender,
                            'row': r
                        })
                        if current_gender == 'M': m_count += 1
                        else: f_count += 1
            else:
                if any(c.isdigit() for c in val0) or val1:
                    print(f"  [G6 Check Row {r} in {sname}]: {row}")
                    
        g6_data.append((sname, section, m_count, f_count, students))

    print("\n=== GRADE 6 SUMMARY ===")
    total_g6 = 0
    for sname, sec, mc, fc, stus in g6_data:
        print(f"Section {sec} ({sname}): {mc} Males, {fc} Females -> Total: {len(stus)}")
        total_g6 += len(stus)
    print(f"TOTAL GRADE 6 STUDENTS: {total_g6}")

parse_all()
