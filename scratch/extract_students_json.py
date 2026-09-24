import openpyxl
import re
import sys
import unicodedata
import json

sys.stdout.reconfigure(encoding='utf-8')

def clean_name(raw_name):
    if not raw_name:
        return ""
    name = unicodedata.normalize('NFC', str(raw_name))
    name = re.sub(r'[\t\r\n]+', ' ', name)
    name = name.strip()
    name = re.sub(r',\s*-\s*$', '', name)
    name = re.sub(r'-\s*$', '', name)
    name = re.sub(r'\s*,\s*', ', ', name)
    name = re.sub(r'\s+', ' ', name)
    return name.strip()

def extract_grade4(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_g4 = []
    
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        section = sheetname.replace('G4', '').replace('g4', '').strip()
        current_gender = 'M'
        sheet_students = []
        
        for r in range(1, ws.max_row + 1):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            non_empty = [c for c in row if c is not None and str(c).strip() != '']
            if not non_empty:
                continue
            r_str = " ".join(str(c).strip() for c in non_empty).upper()
            if any(k in r_str for k in ['PREPARED BY', 'TEACHER', 'ADVISER', 'BCES', 'GRADE 4', 'G4', 'NAME OF']):
                continue
            
            val0 = str(row[0]).strip() if row[0] is not None else ""
            val1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            val2 = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            
            if 'FEMALE' in r_str or 'GIRL' in r_str:
                current_gender = 'F'
                continue
            if 'MALE' in r_str or 'BOY' in r_str:
                current_gender = 'M'
                continue
                
            if val0.isdigit():
                num = int(val0)
                if num == 1 and len(sheet_students) > 0 and current_gender == 'M':
                    current_gender = 'F'
                
                if val1 and not any(k in val1.upper() for k in ['PREPARED', 'TEACHER', 'NAME', 'GRADE', 'ADVISER']):
                    gender = current_gender
                    if val2 in ['M', 'F']:
                        gender = val2
                    cname = clean_name(val1)
                    if cname:
                        sheet_students.append({
                            'grade': '4',
                            'section': section,
                            'sheet': sheetname,
                            'name': cname,
                            'gender': gender,
                            'num_in_sheet': num,
                            'row': r
                        })
        all_g4.extend(sheet_students)
    return all_g4

def extract_grade6(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_g6 = []
    
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        section = sheetname.replace('GRADE 6 -', '').replace('GRADE 6', '').strip()
        current_gender = 'M'
        sheet_students = []
        
        for r in range(1, ws.max_row + 1):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            non_empty = [c for c in row if c is not None and str(c).strip() != '']
            if not non_empty:
                continue
            r_str = " ".join(str(c).strip() for c in non_empty).upper()
            if any(k in r_str for k in ['PREPARED BY', 'TEACHER', 'ADVISER', 'BCES', 'GRADE 6', 'G6', 'SY 202', 'NAME OF LEARNER', 'DATE:']):
                continue
            
            val0 = str(row[0]).strip() if row[0] is not None else ""
            val1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            
            if 'FEMALE' in r_str:
                current_gender = 'F'
                continue
            if 'MALE' in r_str:
                current_gender = 'M'
                continue
                
            if val0.isdigit():
                num = int(val0)
                if val1 and not any(k in val1.upper() for k in ['PREPARED', 'TEACHER', 'NAME', 'GRADE', 'ADVISER']):
                    cname = clean_name(val1)
                    if cname:
                        sheet_students.append({
                            'grade': '6',
                            'section': section,
                            'sheet': sheetname,
                            'name': cname,
                            'gender': current_gender,
                            'num_in_sheet': num,
                            'row': r
                        })
        all_g6.extend(sheet_students)
    return all_g6

g4 = extract_grade4('masterlists/Grade-4-Masterlist-BCES-2026-2027.xlsx')
g6 = extract_grade6('masterlists/GRADE-6-MASTERLIST.xlsx')

print(f"Total Grade 4 students extracted: {len(g4)}")
print(f"Total Grade 6 students extracted: {len(g6)}")

# Check duplicate names within same section
for grade_name, student_list in [('Grade 4', g4), ('Grade 6', g6)]:
    seen = {}
    for s in student_list:
        key = (s['section'], s['name'])
        if key in seen:
            print(f"Duplicate in {grade_name}: {key}")
        seen[key] = s

with open('scratch/extracted_students.json', 'w', encoding='utf-8') as f:
    json.dump({'grade4': g4, 'grade6': g6}, f, ensure_ascii=False, indent=2)

print("Saved scratch/extracted_students.json successfully!")
