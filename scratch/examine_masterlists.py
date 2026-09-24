import openpyxl
import os
import sys

# Ensure UTF-8 stdout
sys.stdout.reconfigure(encoding='utf-8')

ml_dir = 'masterlists'
for f in sorted(os.listdir(ml_dir)):
    fpath = os.path.join(ml_dir, f)
    print("=" * 60)
    print(f"FILE: {f}")
    print("=" * 60)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    print("Sheets:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"\n--- Sheet: '{sheetname}' (max_row={ws.max_row}, max_column={ws.max_column}) ---")
        rows = list(ws.iter_rows(values_only=True))
        # Print non-empty rows
        for idx, row in enumerate(rows, 1):
            if any(cell is not None and str(cell).strip() != '' for cell in row):
                filtered_row = [c for c in row if c is not None]
                if idx <= 15 or idx >= len(rows) - 5:
                    print(f"  Row {idx:2d}: {row[:6]}")
                elif idx == 16:
                    print(f"  ... (total non-empty rows: {len([r for r in rows if any(r)])}) ...")
