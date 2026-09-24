import sqlite3
import openpyxl
import os

conn = sqlite3.connect('ilikesci_db.sqlite')
c = conn.cursor()

print("--- STUDENTS BY GRADE ---")
for r in c.execute("SELECT grade, count(*), min(id), max(id) FROM students GROUP BY grade").fetchall():
    print(f"Grade {r[0]}: count={r[1]}, min_id={r[2]}, max_id={r[3]}")

print("\n--- ALL CURRENT STUDENTS ---")
for r in c.execute("SELECT id, name, grade, section, recitations, total_score FROM students ORDER BY grade, section, name").fetchall():
    print(r)

print("\n--- STUDENT_GRADES BY GRADE_LEVEL ---")
for r in c.execute("SELECT grade_level, count(*) FROM student_grades GROUP BY grade_level").fetchall():
    print(r)

print("\n--- RECITATION_RECORDS BY STUDENT_ID GRADE ---")
for r in c.execute("""
    SELECT s.grade, count(r.id) 
    FROM recitation_records r 
    JOIN students s ON r.student_id = s.id 
    GROUP BY s.grade
""").fetchall():
    print(r)

# Inspect the masterlists
print("\n--- MASTERLISTS INSPECTION ---")
ml_dir = 'masterlists'
for f in os.listdir(ml_dir):
    fpath = os.path.join(ml_dir, f)
    print(f"\nFile: {f}")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    print("Sheets:", wb.sheetnames)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"  Sheet '{sheetname}' max_row={ws.max_row}, max_column={ws.max_column}")
        for row_idx in range(1, min(20, ws.max_row + 1)):
            row_vals = [ws.cell(row=row_idx, column=c).value for c in range(1, min(15, ws.max_column + 1))]
            if any(row_vals):
                print(f"    Row {row_idx}: {row_vals}")
