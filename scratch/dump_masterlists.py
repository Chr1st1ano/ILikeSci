import openpyxl
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')

def parse_g4():
    wb = openpyxl.load_workbook('masterlists/Grade-4-Masterlist-BCES-2026-2027.xlsx', data_only=True)
    print("=== G4 MASTERLIST SHEETS ===")
    total_g4 = 0
    g4_data = {}
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"\nSheet: '{sheetname}', max_row={ws.max_row}, max_col={ws.max_column}")
        students = []
        current_gender = 'M'
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            # Check row
            non_empty = [c for c in row_vals if c is not None and str(c).strip() != '']
            if not non_empty:
                continue
            # Check if this row indicates gender or header or teacher
            row_str = " ".join(str(c).strip() for c in non_empty).upper()
            if any(k in row_str for k in ['GRADE 4', 'G4', 'PREPARED BY', 'TEACHER', 'ADVISER', 'BCES']):
                print(f"  [Header/Footer Row {r}]: {row_vals}")
                continue
            if 'FEMALE' in row_str or 'GIRL' in row_str:
                current_gender = 'F'
                print(f"  [Gender Switch -> F Row {r}]: {row_vals}")
                continue
            if 'MALE' in row_str or 'BOY' in row_str:
                current_gender = 'M'
                print(f"  [Gender Switch -> M Row {r}]: {row_vals}")
                continue
            print(f"  Row {r:2d}: {row_vals}")
    return

def parse_g6():
    wb = openpyxl.load_workbook('masterlists/GRADE-6-MASTERLIST.xlsx', data_only=True)
    print("\n=== G6 MASTERLIST SHEETS ===")
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        print(f"\nSheet: '{sheetname}', max_row={ws.max_row}, max_col={ws.max_column}")
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            non_empty = [c for c in row_vals if c is not None and str(c).strip() != '']
            if not non_empty:
                continue
            row_str = " ".join(str(c).strip() for c in non_empty).upper()
            if any(k in row_str for k in ['GRADE 6', 'G6', 'PREPARED BY', 'TEACHER', 'ADVISER']):
                print(f"  [Header/Footer Row {r}]: {row_vals}")
                continue
            if 'FEMALE' in row_str or 'GIRL' in row_str:
                print(f"  [Gender Switch -> F Row {r}]: {row_vals}")
                continue
            if 'MALE' in row_str or 'BOY' in row_str:
                print(f"  [Gender Switch -> M Row {r}]: {row_vals}")
                continue
            print(f"  Row {r:2d}: {row_vals}")

parse_g4()
parse_g6()
