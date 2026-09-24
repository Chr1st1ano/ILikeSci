import openpyxl

wb = openpyxl.load_workbook('GRADE-4-6_SCIENCE.xlsx', data_only=True)
print("GRADE-4-6_SCIENCE.xlsx sheets:", wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print(f"Sheet '{s}': max_row={ws.max_row}, max_col={ws.max_column}")
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(10, ws.max_column + 1))]
        if any(row_vals):
            print(f"  Row {r}: {row_vals}")
